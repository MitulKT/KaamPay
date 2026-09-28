"""Parses the manufacturer's existing Google-Form/Excel workbook into clean records + a validation report.

Pure (no DB). Handles the layout in the sample: bilingual headers like "कारीगर का नाम (Name of Worker)",
a 'Rate Card' sheet holding a colour-wise quantity block and a lot x work-type rate block side by side,
a 'Submissions' sheet with comma-separated work types/colours, and a 'Payment Details' sheet.
"""
import difflib
import re
from collections import defaultdict
from datetime import datetime

from openpyxl import load_workbook

from .utils import split_bilingual

COLOUR_RE = re.compile(r"^[A-Z]$")


def _en(v) -> str:
    """English key of a possibly bilingual cell: 'फ्रंट (FRONT)' -> 'FRONT'."""
    if v is None:
        return ""
    en, _ = split_bilingual(str(v))
    return re.sub(r"\s+", " ", en).strip().upper()


def _lot(v) -> str | None:
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s[:-2] if s.endswith(".0") else s


def _num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _find_header(ws, predicate, max_rows=6):
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1):
        cells = [_en(c) for c in row]
        res = predicate(cells)
        if res is not None:
            return r_idx, cells, res
    return None


def parse_workbook(path_or_file) -> dict:
    wb = load_workbook(path_or_file, data_only=True, read_only=True)
    data = {"workers": {}, "work_types": {}, "lots": set(), "colour_qty": {}, "rates": {}, "entries": [],
            "payments": []}
    warnings, errors = [], []
    entry_sheets: dict[str, dict] = {}

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head_rows = [[_en(c) for c in r] for r in rows[:6]]

        # --- colour-wise quantity block: header 'LOT' followed by single letters A.. and 'TOTAL'
        for hr_i, cells in enumerate(head_rows):
            for c_i, c in enumerate(cells):
                if c == "LOT" and c_i + 1 < len(cells) and cells[c_i + 1] == "A":
                    cols = {}
                    k = c_i + 1
                    while k < len(cells) and COLOUR_RE.match(cells[k] or ""):
                        cols[k] = cells[k]
                        k += 1
                    total_col = k if k < len(cells) and cells[k] == "TOTAL" else None
                    for r in rows[hr_i + 1:]:
                        lot = _lot(r[c_i] if c_i < len(r) else None)
                        if not lot:
                            continue
                        qty = {code: _num(r[ci]) for ci, code in cols.items() if ci < len(r) and _num(r[ci])}
                        if not qty:
                            continue
                        if lot in data["colour_qty"] and data["colour_qty"][lot] != qty:
                            warnings.append(f"Lot {lot}: colour quantities differ between sheets; using '{ws.title}'")
                        data["colour_qty"][lot] = qty
                        data["lots"].add(lot)
                        if total_col is not None and total_col < len(r) and _num(r[total_col]) is not None:
                            if abs(sum(qty.values()) - _num(r[total_col])) > 0.01:
                                warnings.append(f"Lot {lot}: colour total {sum(qty.values()):g} != sheet Total "
                                                f"{_num(r[total_col]):g} ({ws.title})")

        # --- rate table: header 'LOT NO' followed by work-type columns (skip if next col is 'WORK TYPE')
        for hr_i, cells in enumerate(head_rows):
            for c_i, c in enumerate(cells):
                if c == "LOT NO" and c_i + 1 < len(cells) and cells[c_i + 1] not in ("", "WORK TYPE", "RATE"):
                    cols = {}
                    raw = rows[hr_i]
                    for k in range(c_i + 1, len(cells)):
                        if cells[k] in ("", "TOTAL"):
                            if cells[k] == "TOTAL":
                                break
                            continue
                        cols[k] = cells[k]
                        en, local = split_bilingual(str(raw[k]))
                        data["work_types"].setdefault(cells[k], local)
                    for r in rows[hr_i + 1:]:
                        lot = _lot(r[c_i] if c_i < len(r) else None)
                        if not lot:
                            continue
                        for k, wt in cols.items():
                            v = _num(r[k]) if k < len(r) else None
                            if v is not None:
                                data["rates"][(lot, wt)] = v
                        data["lots"].add(lot)

        # --- single-column masters
        first = head_rows[0] if head_rows else []
        if first and first[0] in ("WORKER NAME", "NAME OF WORKER") and len([c for c in first if c]) == 1:
            for r in rows[1:]:
                if r and r[0]:
                    en, local = split_bilingual(str(r[0]))
                    data["workers"][en.upper()] = {"name": en.title(), "name_local": local}
        if first and first[0] == "LOT NUMBER" and len([c for c in first if c]) == 1:
            for r in rows[1:]:
                lot = _lot(r[0] if r else None)
                if lot:
                    data["lots"].add(lot)
        if first and first[0] == "WORK TYPE" and len([c for c in first[:1] if c]) == 1 and "LOT NO" not in first:
            for r in rows[1:]:
                if r and r[0]:
                    en, local = split_bilingual(str(r[0]))
                    data["work_types"].setdefault(en.upper(), local)

        # --- production entries. Either a Google-Form 'Submissions' sheet (has Timestamp) or a plain
        # 'Production' sheet (Completion Date, no Timestamp). Collected per sheet; picked after the loop so a
        # workbook holding both (the Production tab is usually a copy) is never counted twice.
        if "NAME OF WORKER" in first and "LOT NUMBER" in first and "WORK TYPE" in first:
            has_ts = "TIMESTAMP" in first
            ix = {name: first.index(name) for name in ("NAME OF WORKER", "LOT NUMBER", "WORK TYPE")}
            ix["TIMESTAMP"] = first.index("TIMESTAMP") if has_ts else None
            sheet_entries = entry_sheets.setdefault(ws.title, {"has_ts": has_ts, "entries": []})["entries"]
            colour_i = next((i for i, c in enumerate(first) if c.startswith("COLOR") or c.startswith("COLOUR")), None)
            month_i = next((i for i, c in enumerate(first) if "COMPLETION" in c or c == "MONTH"), None)
            for n, r in enumerate(rows[1:], start=2):
                if not r or not r[ix["NAME OF WORKER"]]:
                    continue
                en, local = split_bilingual(str(r[ix["NAME OF WORKER"]]))
                data["workers"].setdefault(en.upper(), {"name": en.title(), "name_local": local})
                wts = []
                for part in str(r[ix["WORK TYPE"]] or "").split(","):
                    if part.strip():
                        wen, wloc = split_bilingual(part.strip())
                        wts.append(re.sub(r"\s+", " ", wen).upper())
                        data["work_types"].setdefault(wts[-1], wloc)
                colours = [c.strip().upper() for c in str(r[colour_i] if colour_i is not None else "ALL").split(",")
                           if c.strip()]
                colours = ["ALL"] if "ALL" in colours else colours
                ts = r[ix["TIMESTAMP"]] if ix["TIMESTAMP"] is not None else None
                month = r[month_i] if month_i is not None else None
                sheet_entries.append({"row": n, "sheet": ws.title, "worker": en.upper(),
                                        "lot": _lot(r[ix["LOT NUMBER"]]), "work_types": wts, "colours": colours,
                                        "timestamp": ts if isinstance(ts, datetime) else None,
                                        "month": month if isinstance(month, datetime) else None})

        # --- payments
        if first[:4] == ["MONTH", "NAME OF WORKER", "AMOUNT", "DATE"]:
            for n, r in enumerate(rows[1:], start=2):
                if r and r[1] and _num(r[2]):
                    data["payments"].append({"row": n, "name": str(r[1]).strip(), "amount": _num(r[2]),
                                             "date": r[3] if isinstance(r[3], datetime) else None,
                                             "month": r[0] if isinstance(r[0], datetime) else None})

    wb.close()
    # Prefer Form 'Submissions' sheets (with Timestamp); fall back to plain 'Production' sheets.
    chosen = [n for n, v in entry_sheets.items() if v["has_ts"]] or list(entry_sheets)
    for n in chosen:
        data["entries"].extend(entry_sheets[n]["entries"])
    skipped = [n for n in entry_sheets if n not in chosen and entry_sheets[n]["entries"]]
    if skipped:
        warnings.append(f"Used production entries from {', '.join(chosen)}; ignored {', '.join(skipped)} "
                        "(looks like a copy of the same data)")
    return build_report(data, warnings, errors)


def match_name(name: str, worker_keys: list[str]) -> tuple[str | None, float]:
    en = _en(name)
    if en in worker_keys:
        return en, 1.0
    best = difflib.get_close_matches(en, worker_keys, n=1, cutoff=0.6)
    if best:
        return best[0], round(difflib.SequenceMatcher(None, en, best[0]).ratio(), 2)
    return None, 0.0


def build_report(data: dict, warnings: list, errors: list) -> dict:
    """Explodes entries to job lines, prices them, finds duplicates and gaps."""
    lines = []
    for e in data["entries"]:
        if not e["lot"]:
            errors.append(f"{e['sheet']} row {e['row']}: lot number missing")
            continue
        data["lots"].add(e["lot"])
        for wt in e["work_types"]:
            qty = data["colour_qty"].get(e["lot"])
            flags = []
            if qty is None:
                pieces = 0.0
                flags.append("NO_COLOUR_QTY")
            elif e["colours"] == ["ALL"]:
                pieces = sum(qty.values())
            else:
                bad = [c for c in e["colours"] if c not in qty]
                if bad:
                    flags.append("UNKNOWN_COLOUR:" + ",".join(bad))
                pieces = sum(qty.get(c, 0) for c in e["colours"])
            rate = data["rates"].get((e["lot"], wt))
            if rate is None:
                flags.append("NO_RATE")
            lines.append({**e, "work_type": wt, "pieces": pieces, "rate": rate,
                          "amount": round(pieces * (rate or 0), 2), "flags": flags})
    # duplicates: same lot + work type + colour cell claimed more than once
    cells = defaultdict(list)
    for i, l in enumerate(lines):
        codes = sorted(data["colour_qty"].get(l["lot"], {}).keys()) if l["colours"] == ["ALL"] else l["colours"]
        for c in codes:
            cells[(l["lot"], l["work_type"], c)].append(i)
    dup_groups = []
    seen = set()
    for key, idxs in cells.items():
        if len(idxs) > 1:
            g = tuple(sorted(set(idxs)))
            if len(g) > 1 and g not in seen:
                seen.add(g)
                dup_groups.append({"lot": key[0], "work_type": key[1], "colour": key[2], "line_indexes": list(g),
                                   "workers": sorted({lines[i]["worker"] for i in g}),
                                   "same_worker": len({lines[i]["worker"] for i in g}) == 1})
    for i, l in enumerate(lines):
        l["index"] = i
    worker_keys = list(data["workers"].keys())
    name_matches = []
    for p in data["payments"]:
        m, score = match_name(p["name"], worker_keys)
        name_matches.append({"payment_name": p["name"], "proposed_worker": m, "score": score})
    by_month = defaultdict(float)
    for l in lines:
        if l["month"]:
            by_month[l["month"].strftime("%Y-%m")] += l["amount"]
    missing_rates = sorted({(l["lot"], l["work_type"]) for l in lines if l["rate"] is None})
    missing_qty = sorted({l["lot"] for l in lines if "NO_COLOUR_QTY" in l["flags"]})
    return {
        "workers": [{"key": k, **v} for k, v in sorted(data["workers"].items())],
        "work_types": [{"code": k, "name_local": v} for k, v in data["work_types"].items()],
        "lots": sorted(data["lots"], key=lambda x: (len(x), x)),
        "colour_qty": data["colour_qty"],
        "rates": [{"lot": k[0], "work_type": k[1], "rate": v} for k, v in data["rates"].items()],
        "job_lines": lines,
        "duplicates": dup_groups,
        "payments": data["payments"],
        "name_matches": {m["payment_name"]: m for m in name_matches},
        "summary": {
            "workers": len(data["workers"]), "work_types": len(data["work_types"]), "lots": len(data["lots"]),
            "rates": len(data["rates"]), "entries": len(data["entries"]), "job_lines": len(lines),
            "duplicate_groups": len(dup_groups),
            "duplicate_lines": len({i for g in dup_groups for i in g["line_indexes"]}),
            "cross_worker_duplicates": sum(1 for g in dup_groups if not g["same_worker"]),
            "lines_missing_rate": sum(1 for l in lines if l["rate"] is None),
            "lines_missing_qty": sum(1 for l in lines if "NO_COLOUR_QTY" in l["flags"]),
            "payments": len(data["payments"]),
            "amount_by_month": {k: round(v, 2) for k, v in sorted(by_month.items())},
        },
        "missing_rates": [{"lot": a, "work_type": b} for a, b in missing_rates],
        "missing_colour_qty": missing_qty,
        "warnings": warnings,
        "errors": errors,
    }
