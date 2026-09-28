"""Advances, deductions, payout cycles, payments, exports and settlement slips."""
import csv
import html
import io
from datetime import date as _date, timedelta

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font

from ..audit import audit, notify, users_with_role
from ..db import get_db
from ..payout import CycleLine, build_cycle_lines
from ..schemas import AdvanceIn, CycleIn, CycleLinePatch, DeductionIn, PayAllIn, PaymentIn
from ..security import ADMIN, SUPERVISOR, admin_only, anyone, has, staff
from ..utils import as_dt, inr, money, new_id, next_number, now, out, outs, to_ist
from .jobs import company_settings, enrich
from .ledger import worker_summary

router = APIRouter(tags=["money"])


async def _worker(db, company_id, worker_id):
    w = await db.users.find_one({"_id": worker_id, "company_id": company_id})
    if not w:
        raise HTTPException(404, "Worker not found")
    return w


# ---------------- advances ----------------
@router.post("/advances")
async def give_advance(body: AdvanceIn, user: dict = Depends(admin_only)):
    db = get_db()
    w = await _worker(db, user["company_id"], body.worker_id)
    before = await worker_summary(db, user["company_id"], body.worker_id)
    n = await next_number(db, user["company_id"], "advance")
    doc = {"_id": new_id(), "company_id": user["company_id"], "worker_id": body.worker_id, "amount": money(body.amount),
           "date": as_dt(body.date), "mode": body.mode, "reason": body.reason, "recovered_amount": 0.0,
           "status": "OPEN", "receipt_no": f"ADV-{n:05d}", "is_active": True, "created_at": now(),
           "created_by": user["_id"]}
    await db.advances.insert_one(doc)
    await audit(db, user, "advance", doc["_id"], "CREATE", new={"worker": w["name"], "amount": doc["amount"]})
    s = await company_settings(db, user["company_id"])
    if doc["amount"] > s["advance_alert_above"]:
        await notify(db, user["company_id"], await users_with_role(db, user["company_id"], ADMIN), "ADVANCE_ALERT",
                     "Large advance", f"{w['name']} got advance ₹{doc['amount']:.2f}", {"advance_id": doc["_id"]})
    await notify(db, user["company_id"], [body.worker_id], "ADVANCE", "Advance diya / Advance given",
                 f"₹{doc['amount']:.2f} ({body.mode})", {})
    return {"advance": out(doc), "receivable_before": before["receivable"],
            "receivable_after": money(before["receivable"] - doc["amount"])}


@router.get("/advances")
async def list_advances(worker_id: str | None = None, open_only: bool = False, user: dict = Depends(staff)):
    f = {"company_id": user["company_id"], "is_active": True}
    if worker_id:
        f["worker_id"] = worker_id
    if open_only:
        f["status"] = "OPEN"
    return outs([a async for a in get_db().advances.find(f).sort("date", -1)])


@router.delete("/advances/{aid}")
async def void_advance(aid: str, reason: str, user: dict = Depends(admin_only)):
    db = get_db()
    a = await db.advances.find_one({"_id": aid, "company_id": user["company_id"], "is_active": True})
    if not a:
        raise HTTPException(404, "Advance not found")
    if a["recovered_amount"] > 0:
        raise HTTPException(409, "Part of this advance was already recovered in a payout; add a deduction instead")
    await db.advances.update_one({"_id": aid}, {"$set": {"is_active": False, "void_reason": reason}})
    await audit(db, user, "advance", aid, "VOID", old={"amount": a["amount"]}, new={"reason": reason})
    return {"ok": True}


# ---------------- deductions ----------------
@router.post("/deductions")
async def add_deduction(body: DeductionIn, user: dict = Depends(admin_only)):
    db = get_db()
    w = await _worker(db, user["company_id"], body.worker_id)
    doc = {"_id": new_id(), "company_id": user["company_id"], "worker_id": body.worker_id, "amount": money(body.amount),
           "date": as_dt(body.date), "reason": body.reason, "note": body.note, "job_id": body.job_id,
           "payout_cycle_id": None, "is_active": True, "created_at": now(), "created_by": user["_id"]}
    await db.deductions.insert_one(doc)
    await audit(db, user, "deduction", doc["_id"], "CREATE", new={"worker": w["name"], "amount": doc["amount"],
                                                                  "reason": body.reason})
    return out(doc)


@router.get("/deductions")
async def list_deductions(worker_id: str | None = None, user: dict = Depends(staff)):
    f = {"company_id": user["company_id"], "is_active": True}
    if worker_id:
        f["worker_id"] = worker_id
    return outs([d async for d in get_db().deductions.find(f).sort("date", -1)])


# ---------------- payout cycles ----------------
def _line_dict(l: CycleLine) -> dict:
    return {"worker_id": l.worker_id, "gross": l.gross, "advance_recovery": l.advance_recovery,
            "deductions": l.deductions, "net_payable": l.net_payable, "paid_amount": 0.0, "balance": l.net_payable,
            "job_ids": l.job_ids, "deduction_ids": l.deduction_ids, "excluded": False, "recovery_overridden": False}


async def _open_advances(db, company_id) -> dict[str, float]:
    res: dict[str, float] = {}
    async for a in db.advances.find({"company_id": company_id, "is_active": True, "status": "OPEN"}):
        res[a["worker_id"]] = money(res.get(a["worker_id"], 0) + a["amount"] - a["recovered_amount"])
    return res


async def _build_lines(db, company_id: str, from_dt, to_dt, excluded_jobs: set[str]) -> list[dict]:
    s = await company_settings(db, company_id)
    jobs = [j async for j in db.jobs.find({"company_id": company_id, "status": "APPROVED",
                                           "payout_cycle_id": None, "approved_at": {"$gte": from_dt, "$lt": to_dt},
                                           "_id": {"$nin": list(excluded_jobs)}})]
    deds = [d async for d in db.deductions.find({"company_id": company_id, "is_active": True, "payout_cycle_id": None,
                                                 "date": {"$lt": to_dt}})]
    lines = build_cycle_lines(jobs, deds, await _open_advances(db, company_id), s["max_advance_recovery_percent"])
    return [_line_dict(l) for l in lines]


def _totals(lines: list[dict]) -> dict:
    act = [l for l in lines if not l["excluded"]]
    return {k: money(sum(l[k] for l in act)) for k in
            ("gross", "advance_recovery", "deductions", "net_payable", "paid_amount", "balance")} | {"workers": len(act)}


async def _cycle_out(db, c: dict) -> dict:
    d = out(c)
    names = {u["_id"]: u async for u in db.users.find({"_id": {"$in": [l["worker_id"] for l in c["lines"]]}},
                                                      {"name": 1, "name_local": 1})}
    profiles = {p["user_id"]: p async for p in db.worker_profiles.find(
        {"user_id": {"$in": [l["worker_id"] for l in c["lines"]]}})}
    for l in d["lines"]:
        u, p = names.get(l["worker_id"], {}), profiles.get(l["worker_id"], {})
        l["worker_name"], l["worker_name_local"] = u.get("name"), u.get("name_local")
        l["worker_code"], l["payment_mode"], l["upi_id"] = p.get("worker_code"), p.get("payment_mode"), p.get("upi_id")
        l["jobs_count"] = len(l["job_ids"])
    d["totals"] = _totals(c["lines"])
    return d


async def _get_cycle(db, company_id, cycle_id) -> dict:
    c = await db.payout_cycles.find_one({"_id": cycle_id, "company_id": company_id})
    if not c:
        raise HTTPException(404, "Payout cycle not found")
    return c


@router.post("/payouts/cycles")
async def create_cycle(body: CycleIn, user: dict = Depends(admin_only)):
    db = get_db()
    if body.to_date < body.from_date:
        raise HTTPException(400, "To date is before from date")
    from_dt, to_dt = as_dt(body.from_date), as_dt(body.to_date) + timedelta(days=1)
    lines = await _build_lines(db, user["company_id"], from_dt, to_dt, set())
    n = await next_number(db, user["company_id"], "cycle")
    doc = {"_id": new_id(), "company_id": user["company_id"], "cycle_no": f"PC-{n:04d}", "period_type": body.period_type,
           "from_date": from_dt, "to_date": as_dt(body.to_date), "status": "DRAFT", "lines": lines,
           "excluded_job_ids": [], "created_by": user["_id"], "created_at": now()}
    await db.payout_cycles.insert_one(doc)
    await audit(db, user, "payout_cycle", doc["_id"], "CREATE", new={"from": str(body.from_date), "to": str(body.to_date),
                                                                      **_totals(lines)})
    return await _cycle_out(db, doc)


@router.get("/payouts/suggest-period")
async def suggest_period(user: dict = Depends(admin_only)):
    s = await company_settings(get_db(), user["company_id"])
    today = to_ist(now()).date()
    if s["payout_cycle"] == "WEEKLY":
        start = today - timedelta(days=(today.weekday() - s["week_start_day"]) % 7)
        return {"from_date": start, "to_date": start + timedelta(days=6), "period_type": "WEEKLY"}
    start = today.replace(day=1)
    nxt = (start + timedelta(days=32)).replace(day=1)
    return {"from_date": start, "to_date": nxt - timedelta(days=1), "period_type": "MONTHLY"}


@router.get("/payouts/cycles")
async def list_cycles(user: dict = Depends(admin_only)):
    db = get_db()
    res = []
    async for c in db.payout_cycles.find({"company_id": user["company_id"]}).sort("created_at", -1):
        d = out(c)
        d["totals"] = _totals(c["lines"])
        d.pop("lines")
        res.append(d)
    return res


@router.get("/payouts/cycles/{cycle_id}")
async def get_cycle(cycle_id: str, user: dict = Depends(admin_only)):
    db = get_db()
    return await _cycle_out(db, await _get_cycle(db, user["company_id"], cycle_id))


@router.get("/payouts/cycles/{cycle_id}/lines/{worker_id}/jobs")
async def line_jobs(cycle_id: str, worker_id: str, user: dict = Depends(anyone)):
    db = get_db()
    if worker_id != user["_id"] and not has(user, ADMIN):
        raise HTTPException(403, "Not allowed")
    c = await _get_cycle(db, user["company_id"], cycle_id)
    line = next((l for l in c["lines"] if l["worker_id"] == worker_id), None)
    if not line:
        raise HTTPException(404, "Worker not in this cycle")
    jobs = [j async for j in db.jobs.find({"_id": {"$in": line["job_ids"]}})]
    return await enrich(db, user["company_id"], jobs)


@router.post("/payouts/cycles/{cycle_id}/refresh")
async def refresh_cycle(cycle_id: str, user: dict = Depends(admin_only)):
    """Rebuild a DRAFT (e.g. after approving more jobs). Keeps excluded jobs/workers and recovery overrides."""
    db = get_db()
    c = await _get_cycle(db, user["company_id"], cycle_id)
    if c["status"] != "DRAFT":
        raise HTTPException(409, "Only a DRAFT cycle can be refreshed")
    old = {l["worker_id"]: l for l in c["lines"]}
    lines = await _build_lines(db, user["company_id"], c["from_date"], c["to_date"] + timedelta(days=1),
                               set(c.get("excluded_job_ids", [])))
    for l in lines:
        o = old.get(l["worker_id"])
        if o:
            l["excluded"] = o["excluded"]
            if o.get("recovery_overridden"):
                l["advance_recovery"] = min(o["advance_recovery"], l["gross"])
                l["recovery_overridden"] = True
                l["net_payable"] = l["balance"] = money(l["gross"] - l["advance_recovery"] - l["deductions"])
    await db.payout_cycles.update_one({"_id": cycle_id}, {"$set": {"lines": lines}})
    return await get_cycle(cycle_id, user)


@router.patch("/payouts/cycles/{cycle_id}/lines/{worker_id}")
async def patch_line(cycle_id: str, worker_id: str, body: CycleLinePatch, user: dict = Depends(admin_only)):
    db = get_db()
    c = await _get_cycle(db, user["company_id"], cycle_id)
    if c["status"] != "DRAFT":
        raise HTTPException(409, "Cycle is locked")
    if body.exclude_job_ids:
        excluded = set(c.get("excluded_job_ids", [])) | set(body.exclude_job_ids)
        await db.payout_cycles.update_one({"_id": cycle_id}, {"$set": {"excluded_job_ids": list(excluded)}})
        await refresh_cycle(cycle_id, user)
        c = await _get_cycle(db, user["company_id"], cycle_id)
    lines = c["lines"]
    line = next((l for l in lines if l["worker_id"] == worker_id), None)
    if not line:
        raise HTTPException(404, "Worker not in this cycle")
    if body.excluded is not None:
        line["excluded"] = body.excluded
    if body.advance_recovery is not None:
        open_adv = (await _open_advances(db, user["company_id"])).get(worker_id, 0.0)
        if body.advance_recovery > open_adv + 0.001:
            raise HTTPException(400, f"Open advance is only ₹{open_adv:.2f}")
        if body.advance_recovery > line["gross"] - line["deductions"] + 0.001:
            raise HTTPException(400, "Recovery can't be more than the earnings")
        line["advance_recovery"] = money(body.advance_recovery)
        line["recovery_overridden"] = True
    line["net_payable"] = line["balance"] = money(line["gross"] - line["advance_recovery"] - line["deductions"])
    await db.payout_cycles.update_one({"_id": cycle_id}, {"$set": {"lines": lines}})
    await audit(db, user, "payout_cycle", cycle_id, "EDIT_LINE", new={"worker_id": worker_id, **body.model_dump()})
    return await get_cycle(cycle_id, user)


@router.delete("/payouts/cycles/{cycle_id}")
async def delete_draft(cycle_id: str, user: dict = Depends(admin_only)):
    db = get_db()
    c = await _get_cycle(db, user["company_id"], cycle_id)
    if c["status"] != "DRAFT":
        raise HTTPException(409, "Only a DRAFT can be deleted")
    await db.payout_cycles.delete_one({"_id": cycle_id})
    await audit(db, user, "payout_cycle", cycle_id, "DELETE_DRAFT")
    return {"ok": True}


@router.post("/payouts/cycles/{cycle_id}/lock")
async def lock_cycle(cycle_id: str, user: dict = Depends(admin_only)):
    """Freezes the cycle: jobs/deductions get stamped with the cycle, advance recovery is applied (oldest first)."""
    db = get_db()
    c = await _get_cycle(db, user["company_id"], cycle_id)
    if c["status"] != "DRAFT":
        raise HTTPException(409, "Cycle already locked")
    lines = [l for l in c["lines"] if not l["excluded"] and (l["job_ids"] or l["deduction_ids"])]
    # guard: a job may have gone into another cycle meanwhile
    for l in lines:
        taken = await db.jobs.count_documents({"_id": {"$in": l["job_ids"]}, "payout_cycle_id": {"$ne": None}})
        if taken:
            raise HTTPException(409, "Some jobs are already in another cycle. Refresh this draft first.")
    for l in lines:
        await db.jobs.update_many({"_id": {"$in": l["job_ids"]}}, {"$set": {"payout_cycle_id": cycle_id}})
        await db.deductions.update_many({"_id": {"$in": l["deduction_ids"]}}, {"$set": {"payout_cycle_id": cycle_id}})
        left = l["advance_recovery"]
        async for a in db.advances.find({"company_id": user["company_id"], "worker_id": l["worker_id"],
                                         "is_active": True, "status": "OPEN"}).sort("date", 1):
            if left <= 0:
                break
            take = min(left, money(a["amount"] - a["recovered_amount"]))
            rec = money(a["recovered_amount"] + take)
            await db.advances.update_one({"_id": a["_id"]}, {"$set": {
                "recovered_amount": rec, "status": "RECOVERED" if rec >= a["amount"] - 0.001 else "OPEN"}})
            left = money(left - take)
        if l["net_payable"] <= 0.001:  # fully settled by advance recovery / deductions: nothing to pay
            await db.jobs.update_many({"_id": {"$in": l["job_ids"]}, "status": "APPROVED"},
                                      {"$set": {"status": "PAID", "paid_at": now()}})
    if all(l["net_payable"] <= 0.001 for l in lines):
        await db.payout_cycles.update_one({"_id": cycle_id}, {"$set": {"status": "PAID", "lines": lines,
                                                                       "locked_by": user["_id"], "locked_at": now()}})
        await audit(db, user, "payout_cycle", cycle_id, "LOCK", new=_totals(lines))
        return await get_cycle(cycle_id, user)
    await db.payout_cycles.update_one({"_id": cycle_id}, {"$set": {"status": "LOCKED", "lines": lines,
                                                                   "locked_by": user["_id"], "locked_at": now()}})
    await audit(db, user, "payout_cycle", cycle_id, "LOCK", new=_totals(lines))
    return await get_cycle(cycle_id, user)


# ---------------- payments ----------------
async def _record_payment(db, user: dict, body: PaymentIn) -> dict:
    w = await _worker(db, user["company_id"], body.worker_id)
    pay_dt = as_dt(body.date)
    if not body.confirm_duplicate:
        day = pay_dt.replace(hour=0, minute=0, second=0, microsecond=0)
        dup = await db.payments.find_one({"company_id": user["company_id"], "worker_id": body.worker_id,
                                          "amount": money(body.amount), "is_active": True,
                                          "date": {"$gte": day, "$lt": day + timedelta(days=1)}})
        if dup:
            raise HTTPException(409, {"message": f"Same amount already paid to {w['name']} on this day "
                                                 f"(receipt {dup['receipt_no']}). Send confirm_duplicate=true if intended.",
                                      "duplicate_of": dup["receipt_no"]})
    cycle = None
    if body.payout_cycle_id:
        cycle = await _get_cycle(db, user["company_id"], body.payout_cycle_id)
        if cycle["status"] == "DRAFT":
            raise HTTPException(409, "Lock the cycle before paying")
        line = next((l for l in cycle["lines"] if l["worker_id"] == body.worker_id), None)
        if not line:
            raise HTTPException(400, "Worker is not in this cycle")
        if body.amount > line["balance"] + 0.001:
            raise HTTPException(400, f"Balance for this cycle is only ₹{line['balance']:.2f}")
    n = await next_number(db, user["company_id"], "receipt")
    doc = {"_id": new_id(), "company_id": user["company_id"], "worker_id": body.worker_id,
           "payout_cycle_id": body.payout_cycle_id, "amount": money(body.amount), "date": pay_dt, "mode": body.mode,
           "reference_no": body.reference_no, "paid_by": user["_id"], "receipt_no": f"RCP-{n:05d}",
           "is_active": True, "created_at": now()}
    await db.payments.insert_one(doc)
    if cycle:
        lines = cycle["lines"]
        for l in lines:
            if l["worker_id"] == body.worker_id:
                l["paid_amount"] = money(l["paid_amount"] + doc["amount"])
                l["balance"] = money(l["net_payable"] - l["paid_amount"])
                if l["balance"] <= 0.001:
                    await db.jobs.update_many({"_id": {"$in": l["job_ids"]}, "status": "APPROVED"},
                                              {"$set": {"status": "PAID", "paid_at": now()},
                                               "$push": {"status_history": {"status": "PAID", "by": user["_id"],
                                                                            "by_name": user["name"], "at": now(),
                                                                            "note": doc["receipt_no"]}}})
        status = "PAID" if all(l["balance"] <= 0.001 for l in lines) else "LOCKED"
        await db.payout_cycles.update_one({"_id": cycle["_id"]}, {"$set": {"lines": lines, "status": status}})
    await audit(db, user, "payment", doc["_id"], "CREATE", new={"worker": w["name"], "amount": doc["amount"],
                                                                "mode": body.mode, "cycle": body.payout_cycle_id})
    await notify(db, user["company_id"], [body.worker_id], "PAID", "Paisa mila / Payment done 💰",
                 f"₹{doc['amount']:.2f} ({body.mode}) receipt {doc['receipt_no']}",
                 {"payment_id": doc["_id"], "cycle_id": body.payout_cycle_id})
    return out(doc)


@router.post("/payments")
async def record_payment(body: PaymentIn, user: dict = Depends(admin_only)):
    return await _record_payment(get_db(), user, body)


@router.post("/payouts/cycles/{cycle_id}/pay-all")
async def pay_all(cycle_id: str, body: PayAllIn, user: dict = Depends(admin_only)):
    db = get_db()
    c = await _get_cycle(db, user["company_id"], cycle_id)
    paid = []
    profiles = {p["user_id"]: p async for p in db.worker_profiles.find({"company_id": user["company_id"]})}
    for l in c["lines"]:
        if l["balance"] > 0.001:
            mode = body.mode if body.mode != "AUTO" else profiles.get(l["worker_id"], {}).get("payment_mode", "CASH")
            p = await _record_payment(db, user, PaymentIn(worker_id=l["worker_id"], amount=l["balance"], mode=mode,
                                                          reference_no=body.reference_no, date=body.date,
                                                          payout_cycle_id=cycle_id, confirm_duplicate=True))
            paid.append(p)
    return {"payments": paid, "total": money(sum(p["amount"] for p in paid))}


@router.get("/payments")
async def list_payments(worker_id: str | None = None, cycle_id: str | None = None, user: dict = Depends(anyone)):
    f = {"company_id": user["company_id"], "is_active": True}
    if not has(user, ADMIN, SUPERVISOR):
        f["worker_id"] = user["_id"]
    elif worker_id:
        f["worker_id"] = worker_id
    if cycle_id:
        f["payout_cycle_id"] = cycle_id
    return outs([p async for p in get_db().payments.find(f).sort("date", -1)])


@router.delete("/payments/{pid}")
async def void_payment(pid: str, reason: str, user: dict = Depends(admin_only)):
    db = get_db()
    p = await db.payments.find_one({"_id": pid, "company_id": user["company_id"], "is_active": True})
    if not p:
        raise HTTPException(404, "Payment not found")
    await db.payments.update_one({"_id": pid}, {"$set": {"is_active": False, "void_reason": reason}})
    if p.get("payout_cycle_id"):
        c = await db.payout_cycles.find_one({"_id": p["payout_cycle_id"]})
        for l in c["lines"]:
            if l["worker_id"] == p["worker_id"]:
                l["paid_amount"] = money(l["paid_amount"] - p["amount"])
                l["balance"] = money(l["net_payable"] - l["paid_amount"])
                await db.jobs.update_many({"_id": {"$in": l["job_ids"]}, "status": "PAID"},
                                          {"$set": {"status": "APPROVED", "paid_at": None}})
        await db.payout_cycles.update_one({"_id": c["_id"]}, {"$set": {"lines": c["lines"], "status": "LOCKED"}})
    await audit(db, user, "payment", pid, "VOID", old={"amount": p["amount"]}, new={"reason": reason})
    return {"ok": True}


# ---------------- exports & slips ----------------
@router.get("/payouts/cycles/{cycle_id}/export.xlsx")
async def export_cycle_xlsx(cycle_id: str, user: dict = Depends(admin_only)):
    db = get_db()
    c = await _cycle_out(db, await _get_cycle(db, user["company_id"], cycle_id))
    wb = Workbook()
    ws = wb.active
    ws.title = "Payout"
    ws.append([f"Payout {c['cycle_no']}  {c['from_date']:%d-%m-%Y} to {c['to_date']:%d-%m-%Y}  ({c['status']})"])
    ws["A1"].font = Font(bold=True, size=13)
    head = ["Code", "Worker", "Jobs", "Gross", "Advance cut", "Deductions", "Net payable", "Paid", "Balance", "Mode",
            "UPI", "Signature"]
    ws.append(head)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for l in c["lines"]:
        ws.append([l["worker_code"], l["worker_name"], l["jobs_count"], l["gross"], l["advance_recovery"],
                   l["deductions"], l["net_payable"], l["paid_amount"], l["balance"], l["payment_mode"], l["upi_id"], ""])
    t = c["totals"]
    ws.append(["", "TOTAL", "", t["gross"], t["advance_recovery"], t["deductions"], t["net_payable"], t["paid_amount"],
               t["balance"]])
    for row in ws.iter_rows(min_row=3, min_col=4, max_col=9):
        for cell in row:
            cell.number_format = "#,##,##0.00"
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["L"].width = 18
    # job-wise sheet
    js = wb.create_sheet("Jobs")
    js.append(["Worker", "Job", "Lot", "Work type", "Colours", "Pieces", "Rate", "Amount", "Approved"])
    for l in c["lines"]:
        for j in await enrich(db, user["company_id"], [x async for x in db.jobs.find({"_id": {"$in": l["job_ids"]}})]):
            js.append([l["worker_name"], j["job_no"], j["lot_no"], j["work_type_code"], ",".join(j["colour_codes"]),
                       j["pieces"], j["rate_snapshot"], j["amount"], to_ist(j["approved_at"]).strftime("%d-%m-%Y")])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={c['cycle_no']}.xlsx"})


@router.get("/payouts/cycles/{cycle_id}/bank.csv")
async def export_bank_csv(cycle_id: str, user: dict = Depends(admin_only)):
    """UPI/bank bulk-pay file: one row per worker with balance > 0 who is paid by UPI or BANK."""
    db = get_db()
    c = await _cycle_out(db, await _get_cycle(db, user["company_id"], cycle_id))
    profiles = {p["user_id"]: p async for p in db.worker_profiles.find({"company_id": user["company_id"]})}
    buf = io.StringIO()
    wr = csv.writer(buf)
    wr.writerow(["worker_code", "name", "mode", "upi_id", "bank_name", "ifsc", "account_last4", "amount", "narration"])
    for l in c["lines"]:
        p = profiles.get(l["worker_id"], {})
        if l["balance"] > 0 and p.get("payment_mode") in ("UPI", "BANK"):
            wr.writerow([p.get("worker_code"), l["worker_name"], p.get("payment_mode"), p.get("upi_id"),
                         p.get("bank_name"), p.get("ifsc"), p.get("account_last4"), f"{l['balance']:.2f}",
                         f"Job work {c['cycle_no']}"])
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": f"attachment; filename={c['cycle_no']}-bank.csv"})


@router.get("/payouts/cycles/{cycle_id}/slip/{worker_id}", response_class=HTMLResponse)
async def settlement_slip(cycle_id: str, worker_id: str, user: dict = Depends(anyone)):
    """Printable HTML slip. The app turns it into a PDF (expo-print) and shares it on WhatsApp."""
    db = get_db()
    if worker_id != user["_id"] and not has(user, ADMIN):
        raise HTTPException(403, "Not allowed")
    c = await _get_cycle(db, user["company_id"], cycle_id)
    line = next((l for l in c["lines"] if l["worker_id"] == worker_id), None)
    if not line:
        raise HTTPException(404, "Worker not in this cycle")
    company = await db.companies.find_one({"_id": user["company_id"]})
    w = await db.users.find_one({"_id": worker_id})
    p = await db.worker_profiles.find_one({"user_id": worker_id}) or {}
    jobs = await enrich(db, user["company_id"], [j async for j in db.jobs.find({"_id": {"$in": line["job_ids"]}})])
    pays = [x async for x in db.payments.find({"payout_cycle_id": cycle_id, "worker_id": worker_id, "is_active": True})]
    e = html.escape
    rows = "".join(
        f"<tr><td>{e(str(j['lot_no']))}</td><td>{e(j['work_type_code'] or '')}</td><td>{e(','.join(j['colour_codes']))}</td>"
        f"<td class=r>{j['pieces']:g}</td><td class=r>{j['rate_snapshot']:.2f}</td><td class=r>{inr(j['amount'])}</td></tr>"
        for j in sorted(jobs, key=lambda j: (str(j['lot_no']), j['work_type_code'] or '')))
    pay_rows = "".join(f"<div>{to_ist(x['date']):%d-%m-%Y} · {x['mode']} · {e(x['receipt_no'])} · <b>{inr(x['amount'])}</b></div>"
                       for x in pays) or "<div>Not paid yet</div>"
    doc = f"""<!doctype html><html><head><meta charset=utf-8><meta name=viewport content="width=device-width">
<title>Slip {e(c['cycle_no'])} {e(w['name'])}</title>
<style>body{{font-family:'Noto Sans','Noto Sans Devanagari',sans-serif;margin:16px;color:#111}}h1{{font-size:18px;margin:0}}
table{{width:100%;border-collapse:collapse;font-size:12px;margin-top:10px}}td,th{{border-bottom:1px solid #ddd;padding:5px;text-align:left}}
.r{{text-align:right}}.box{{display:flex;justify-content:space-between;padding:4px 0}}.net{{font-size:20px;font-weight:700}}
.muted{{color:#666;font-size:12px}}</style></head><body>
<h1>{e(company['name'])}</h1><div class=muted>Job work settlement slip · {e(c['cycle_no'])}</div>
<p><b>{e(w['name'])}</b> {e(w.get('name_local') or '')} · {e(p.get('worker_code') or '')}<br>
<span class=muted>Period {c['from_date']:%d-%m-%Y} to {c['to_date']:%d-%m-%Y}</span></p>
<table><tr><th>Lot</th><th>Work</th><th>Colour</th><th class=r>Pcs</th><th class=r>Rate</th><th class=r>Amount</th></tr>{rows}</table>
<div style="margin-top:12px">
<div class=box><span>Gross earnings</span><b>{inr(line['gross'])}</b></div>
<div class=box><span>Advance recovered</span><span>- {inr(line['advance_recovery'])}</span></div>
<div class=box><span>Deductions</span><span>- {inr(line['deductions'])}</span></div>
<div class="box net"><span>Net payable</span><span>{inr(line['net_payable'])}</span></div>
<div class=box><span>Paid</span><span>{inr(line['paid_amount'])}</span></div>
<div class=box><span>Balance</span><span>{inr(line['balance'])}</span></div></div>
<h3 style="font-size:14px">Payments</h3>{pay_rows}
<p class=muted>Generated {to_ist(now()):%d-%m-%Y %H:%M} · KaamPay</p></body></html>"""
    return HTMLResponse(doc)
