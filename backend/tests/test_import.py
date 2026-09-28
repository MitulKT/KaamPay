"""Imports the real sample workbook the manufacturer shared."""
import pytest

from app.importer import parse_workbook

from .conftest import SAMPLE


def test_parse_sample_workbook():
    rep = parse_workbook(SAMPLE)
    s = rep["summary"]
    print(s)
    assert s["workers"] >= 31
    assert s["work_types"] == 27
    assert s["lots"] >= 37
    assert s["job_lines"] > 1000
    assert s["amount_by_month"]["2026-08"] == 501684.0
    assert s["duplicate_groups"] > 50 and s["cross_worker_duplicates"] >= 5
    assert s["lines_missing_rate"] > 0
    # the worked example: Suraj lot 128 colour A FRONT = 77 x 14.5
    l = next(l for l in rep["job_lines"] if l["worker"] == "SURAJ" and l["lot"] == "128" and l["work_type"] == "FRONT")
    assert l["pieces"] == 77 and l["rate"] == 14.5 and l["amount"] == 1116.5
    assert rep["name_matches"]["Gulam"]["proposed_worker"] == "GULAM"


async def test_import_commit_and_undo(admin, db):
    with open(SAMPLE, "rb") as f:
        r = await admin.post("/import/preview", files={"file": ("sample.xlsx", f.read(),
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    prev = r.json()
    sid = prev["session_id"]
    r = await admin.post(f"/import/{sid}/commit", {"settled_before_month": "2026-08", "duplicate_policy": "KEEP_FIRST"})
    assert r.status_code == 200, r.text
    res = r.json()
    print(res)
    assert res["counts"]["jobs"] > 900 and res["counts"]["lines_skipped"] > 0
    assert res["counts"]["opening_cycles"] >= 3
    # settled months -> PAID with opening settlement; Aug/Sep -> APPROVED, ready for a payout cycle
    assert await db.jobs.count_documents({"status": "PAID"}) > 0
    assert await db.jobs.count_documents({"status": "APPROVED"}) > 0
    # a worker's ledger for a settled month nets to zero; open months show receivable
    d = (await admin.get("/dashboard/admin")).json()
    assert d["kpis"]["approved_unpaid"]["amount"] > 100000
    assert any(a["kind"] == "PENDING_MOBILE" for a in d["alerts"])
    m = (await admin.get("/reports/monthly-summary")).json()
    assert len(m["rows"]) >= 4
    # undo removes everything
    r = await admin.delete(f"/import/{sid}")
    assert r.status_code == 200
    assert await db.jobs.count_documents({}) == 0
    assert await db.lots.count_documents({}) == 0


def test_production_sheet_without_timestamp(tmp_path):
    """A plain 'Production' sheet (Completion Date, no Form Timestamp) must still be read as entries,
    and when a Form 'Submissions' sheet is also present only one of them is used (no double count)."""
    from datetime import datetime

    from openpyxl import Workbook

    from app.importer import parse_workbook

    def build(with_submissions: bool):
        wb = Workbook()
        rc = wb.active
        rc.title = "Rate Card"
        rc.append(["Rate Card", "Colorwise Quantity"])
        rc.append(["Lot No", "FRONT", "", "Lot", "A", "B"])
        rc.append([501, 10, "", 501, 20, 30])
        prod = wb.create_sheet("Production")
        prod.append(["जॉब कम्पलीशन डेट (Work Completion Date)", "कारीगर का नाम (Name of Worker)",
                     "लोट नंबर (Lot Number)", "जॉब का प्रकार (Work Type)", "कलर कोड (Color Code)", "Rate", "Qty", "Total"])
        prod.append([datetime(2026, 9, 1), "सूरज (SURAJ)", 501, "फ्रंट (FRONT)", "A", 10, 20, 200])
        prod.append([datetime(2026, 9, 1), "गुलाम (GULAM)", 501, "फ्रंट (FRONT)", "B", 10, 30, 300])
        if with_submissions:
            sub = wb.create_sheet("Submissions")
            sub.append(["Timestamp", "कारीगर का नाम (Name of Worker)", "लोट नंबर (Lot Number)",
                        "जॉब का प्रकार (Work Type)", "कलर कोड (Color Code)", "जॉब कम्पलीशन डेट (Work Completion Date)"])
            sub.append([datetime(2026, 9, 2, 10), "सूरज (SURAJ)", 501, "फ्रंट (FRONT)", "A", datetime(2026, 9, 1)])
            sub.append([datetime(2026, 9, 2, 11), "गुलाम (GULAM)", 501, "फ्रंट (FRONT)", "B", datetime(2026, 9, 1)])
        path = tmp_path / f"wb_{with_submissions}.xlsx"
        wb.save(path)
        return parse_workbook(str(path))

    only_prod = build(False)
    assert only_prod["summary"]["entries"] == 2
    assert only_prod["summary"]["amount_by_month"] == {"2026-09": 500.0}

    both = build(True)
    assert both["summary"]["entries"] == 2  # Submissions used, Production copy ignored
    assert any("ignored Production" in w for w in both["warnings"])
